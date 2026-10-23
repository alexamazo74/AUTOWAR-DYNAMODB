"""
Technical Report Generator Lambda
Generates detailed technical reports in XLSX format
"""

import json
import logging
import io
import os
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from autowar_common.report_storage import upload_report, generate_signed_url

logger = logging.getLogger()
logger.setLevel(logging.INFO)

# GFourmis corporate palette (openpyxl hex, no '#')
GF_ORANGE = 'F06000'  # primary orange - main headers
GF_RED = 'C01010'     # corporate red - section/question headers


def resolve_asset_path(filename):
    """Resolve a branding asset path with a Lambda-first fallback.

    Tries the Lambda layer mount (/opt/python/autowar_assets/<filename>) first,
    then falls back to the project-relative path (lambda/layers/reports/python/
    autowar_assets/<filename>) computed from this handler's location so it also
    works in local tests. Returns the first existing path or None.
    """
    lambda_path = os.path.join('/opt/python', 'autowar_assets', filename)
    if os.path.exists(lambda_path):
        return lambda_path

    handler_dir = os.path.dirname(os.path.abspath(__file__))
    local_path = os.path.abspath(os.path.join(
        handler_dir, '..', 'layers', 'reports', 'python', 'autowar_assets', filename
    ))
    if os.path.exists(local_path):
        return local_path

    logger.info(f"Branding asset not found: {filename} (tried {lambda_path} and {local_path})")
    return None


def lambda_handler(event, context):
    """
    Generate technical report in XLSX format
    
    Input:
    {
        "accountId": "...",
        "evaluationTimestamp": "...",
        "evaluationData": {...},
        "selectedPilares": ["security", "reliability", ...]
    }
    
    Output:
    {
        "s3Key": "...",
        "signedUrl": "...",
        "expiresIn": 3600
    }
    """
    try:
        # Parse input
        if isinstance(event.get('body'), str):
            body = json.loads(event['body'])
        else:
            body = event
        
        account_id = body.get('accountId')
        evaluation_timestamp = body.get('evaluationTimestamp')
        evaluation_data = body.get('evaluationData')
        selected_pilares = body.get('selectedPilares', [])
        user_info = body.get('userInfo', {})
        
        logger.info(f"Generating technical report for account {account_id}")
        
        # Generate XLSX report
        xlsx_content = generate_technical_report(
            evaluation_data,
            selected_pilares,
            account_id,
            evaluation_timestamp,
            user_info
        )
        
        # Upload to S3
        s3_key = upload_report(
            account_id,
            evaluation_timestamp,
            'technical',
            xlsx_content,
            'xlsx'
        )
        
        # Generate signed URL
        signed_url = generate_signed_url(s3_key, expiration_seconds=3600)
        
        logger.info(f"Technical report generated: {s3_key}")
        
        return {
            'statusCode': 200,
            'headers': {
                'Content-Type': 'application/json',
                'Access-Control-Allow-Origin': '*'
            },
            'body': json.dumps({
                's3Key': s3_key,
                'signedUrl': signed_url,
                'expiresIn': 3600
            })
        }
        
    except Exception as e:
        logger.error(f"Error generating technical report: {str(e)}", exc_info=True)
        return {
            'statusCode': 500,
            'headers': {
                'Content-Type': 'application/json',
                'Access-Control-Allow-Origin': '*'
            },
            'body': json.dumps({
                'error': 'Failed to generate technical report',
                'message': str(e)
            })
        }


def generate_technical_report(evaluation_data, selected_pilares, account_id, evaluation_timestamp, user_info=None):
    """Generate technical report XLSX content"""
    
    wb = Workbook()
    
    # Remove default sheet
    if 'Sheet' in wb.sheetnames:
        wb.remove(wb['Sheet'])
    
    # Create summary sheet
    create_summary_sheet(wb, evaluation_data, selected_pilares, account_id, evaluation_timestamp, user_info)
    
    # Create one sheet per selected pilar
    for pilar in selected_pilares:
        create_pilar_sheet(wb, evaluation_data, pilar, account_id, user_info)
    
    # Save to bytes
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    
    return output.read()


def create_summary_sheet(wb, evaluation_data, selected_pilares, account_id, evaluation_timestamp, user_info=None):
    """Create summary worksheet with detailed statistics and charts"""
    from openpyxl.chart import BarChart, PieChart, Reference
    from openpyxl.drawing.image import Image as XLImage
    
    ws = wb.create_sheet("Resumen", 0)
    
    # GFourmis logo (top-right, over the title area without covering A1:H1 text)
    logo_path = resolve_asset_path('gfourmis-logo.png')
    if logo_path:
        try:
            img = XLImage(logo_path)
            # Keep aspect ratio ~2712x2478
            img.width = 180
            img.height = 164
            ws.add_image(img, 'G1')
        except Exception as e:
            logger.info(f"Could not embed logo in summary sheet: {e}")
    
    # Title
    ws['A1'] = 'Reporte Técnico - AWS Well-Architected Review'
    ws['A1'].font = Font(size=16, bold=True, color=GF_ORANGE)
    ws.merge_cells('A1:F1')
    
    # Metadata
    ws['A3'] = 'Cuenta AWS:'
    ws['B3'] = account_id
    ws['A4'] = 'Fecha de Evaluación:'
    ws['B4'] = evaluation_timestamp
    ws['A5'] = 'Fecha de Generación:'
    ws['B5'] = datetime.utcnow().isoformat() + 'Z'
    
    # User information
    if user_info:
        ws['A6'] = 'Usuario:'
        ws['B6'] = f"{user_info.get('user_name', 'Unknown')} ({user_info.get('user_type', 'Unknown')})"
        ws['A7'] = 'User ID:'
        ws['B7'] = user_info.get('user_id', 'Unknown')
    
    # Pillar scores header
    ws['A9'] = 'Scores por Pilar'
    ws['A9'].font = Font(size=14, bold=True)
    
    # Pillar scores table
    row = 10
    ws[f'A{row}'] = 'Pilar'
    ws[f'B{row}'] = 'Score (%)'
    apply_header_style(ws, row, 2)
    
    pillar_scores = evaluation_data.get('pillarScores', {})
    pillar_names = {
        'security': 'Seguridad',
        'reliability': 'Fiabilidad',
        'performance': 'Rendimiento',
        'cost': 'Optimización de Costos',
        'operationalExcellence': 'Excelencia Operacional',
        'sustainability': 'Sostenibilidad'
    }
    
    row += 1
    score_start_row = row
    for pilar in selected_pilares:
        ws[f'A{row}'] = pillar_names.get(pilar, pilar)
        ws[f'B{row}'] = pillar_scores.get(pilar, 0)
        row += 1
    score_end_row = row - 1
    
    # Detailed statistics by pilar
    row += 2
    ws[f'A{row}'] = 'Estadísticas Detalladas por Pilar'
    ws[f'A{row}'].font = Font(size=14, bold=True)
    row += 1
    
    # Headers for detailed stats
    headers = ['Pilar', 'Total BPs', 'Cumplen', 'No Cumplen', 'Pendientes', 
               'Cumplen-Crítica', 'Cumplen-Alta', 'Cumplen-Media', 'Cumplen-Baja',
               'No Cumplen-Crítica', 'No Cumplen-Alta', 'No Cumplen-Media', 'No Cumplen-Baja',
               'Pendientes-Crítica', 'Pendientes-Alta', 'Pendientes-Media', 'Pendientes-Baja']
    
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col, value=header)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill(start_color=GF_ORANGE, end_color=GF_ORANGE, fill_type='solid')
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    
    row += 1
    stats_start_row = row
    
    # Calculate statistics for each pilar
    best_practices = evaluation_data.get('bestPractices', [])
    
    for pilar in selected_pilares:
        pilar_normalized = pilar.replace('_', '').lower()
        filtered_bps = [
            bp for bp in best_practices 
            if bp.get('pilar', '').replace('_', '').lower() == pilar_normalized
        ]
        
        # Calculate counts
        total_bps = len(filtered_bps)
        compliant = sum(1 for bp in filtered_bps if bp.get('status') == 'COMPLIANT')
        non_compliant = sum(1 for bp in filtered_bps if bp.get('status') == 'NON_COMPLIANT')
        pending = sum(1 for bp in filtered_bps if bp.get('status') == 'PENDING_REVIEW')
        
        # Severity breakdown by status
        def count_by_status_severity(status, severity):
            return sum(1 for bp in filtered_bps 
                      if bp.get('status') == status and bp.get('severity') == severity)
        
        # Write data
        col = 1
        ws.cell(row=row, column=col, value=pillar_names.get(pilar, pilar))
        col += 1
        ws.cell(row=row, column=col, value=total_bps)
        col += 1
        ws.cell(row=row, column=col, value=compliant)
        col += 1
        ws.cell(row=row, column=col, value=non_compliant)
        col += 1
        ws.cell(row=row, column=col, value=pending)
        col += 1
        
        # Compliant by severity
        for severity in ['CRITICAL', 'HIGH', 'MEDIUM', 'LOW']:
            ws.cell(row=row, column=col, value=count_by_status_severity('COMPLIANT', severity))
            col += 1
        
        # Non-compliant by severity
        for severity in ['CRITICAL', 'HIGH', 'MEDIUM', 'LOW']:
            ws.cell(row=row, column=col, value=count_by_status_severity('NON_COMPLIANT', severity))
            col += 1
        
        # Pending by severity
        for severity in ['CRITICAL', 'HIGH', 'MEDIUM', 'LOW']:
            ws.cell(row=row, column=col, value=count_by_status_severity('PENDING_REVIEW', severity))
            col += 1
        
        row += 1
    
    stats_end_row = row - 1
    
    # Add charts
    row += 2
    
    # Chart 1: Pillar Scores (Bar Chart)
    chart1 = BarChart()
    chart1.title = "Scores por Pilar (%)"
    chart1.style = 10
    chart1.y_axis.title = 'Score (%)'
    chart1.x_axis.title = 'Pilar'
    
    data = Reference(ws, min_col=2, min_row=score_start_row-1, max_row=score_end_row)
    cats = Reference(ws, min_col=1, min_row=score_start_row, max_row=score_end_row)
    chart1.add_data(data, titles_from_data=True)
    chart1.set_categories(cats)
    chart1.height = 10
    chart1.width = 20
    
    ws.add_chart(chart1, f'A{row}')
    
    # Chart 2: Status Distribution (Stacked Bar Chart)
    chart2 = BarChart()
    chart2.type = "col"
    chart2.style = 11
    chart2.title = "Distribución de Estados por Pilar"
    chart2.y_axis.title = 'Cantidad de BPs'
    chart2.x_axis.title = 'Pilar'
    chart2.grouping = "stacked"
    chart2.overlap = 100
    
    # Data for Cumplen, No Cumplen, Pendientes (columns 3, 4, 5)
    data = Reference(ws, min_col=3, min_row=stats_start_row-1, max_col=5, max_row=stats_end_row)
    cats = Reference(ws, min_col=1, min_row=stats_start_row, max_row=stats_end_row)
    chart2.add_data(data, titles_from_data=True)
    chart2.set_categories(cats)
    chart2.height = 10
    chart2.width = 20
    
    ws.add_chart(chart2, f'L{row}')
    
    # Chart 3: Overall Status Distribution (Pie Chart)
    row += 18
    chart3 = PieChart()
    chart3.title = "Distribución Global de Estados"
    chart3.style = 10
    
    # Calculate overall totals
    overall_row = row
    ws.cell(row=overall_row, column=1, value='Estado')
    ws.cell(row=overall_row, column=2, value='Cantidad')
    ws.cell(row=overall_row, column=1).font = Font(bold=True)
    ws.cell(row=overall_row, column=2).font = Font(bold=True)
    
    overall_row += 1
    ws.cell(row=overall_row, column=1, value='Cumplen')
    ws.cell(row=overall_row, column=2, value=f'=SUM(C{stats_start_row}:C{stats_end_row})')
    overall_row += 1
    ws.cell(row=overall_row, column=1, value='No Cumplen')
    ws.cell(row=overall_row, column=2, value=f'=SUM(D{stats_start_row}:D{stats_end_row})')
    overall_row += 1
    ws.cell(row=overall_row, column=1, value='Pendientes')
    ws.cell(row=overall_row, column=2, value=f'=SUM(E{stats_start_row}:E{stats_end_row})')
    
    data = Reference(ws, min_col=2, min_row=row, max_row=row+3)
    cats = Reference(ws, min_col=1, min_row=row+1, max_row=row+3)
    chart3.add_data(data, titles_from_data=True)
    chart3.set_categories(cats)
    chart3.height = 10
    chart3.width = 12
    
    ws.add_chart(chart3, f'A{row}')
    
    # Auto-adjust column widths
    ws.column_dimensions['A'].width = 30
    ws.column_dimensions['B'].width = 15
    for col in range(3, 18):
        ws.column_dimensions[get_column_letter(col)].width = 12


def create_pilar_sheet(wb, evaluation_data, pilar, account_id=None, user_info=None):
    """Create worksheet for a specific pilar with grouping by question"""
    
    pilar_names = {
        'security': 'Seguridad',
        'reliability': 'Fiabilidad',
        'performance': 'Rendimiento',
        'cost': 'Optimización de Costos',
        'operationalExcellence': 'Excelencia Operacional',
        'sustainability': 'Sostenibilidad'
    }
    
    sheet_name = pilar_names.get(pilar, pilar)[:31]  # Excel sheet name limit
    ws = wb.create_sheet(sheet_name)
    
    # Add header with account and user info
    current_row = 1
    if account_id or user_info:
        ws.merge_cells(f'A{current_row}:H{current_row}')
        header_cell = ws.cell(row=current_row, column=1)
        header_parts = []
        if account_id:
            header_parts.append(f'Cuenta: {account_id}')
        if user_info and user_info.get('user_name'):
            header_parts.append(f'Usuario: {user_info.get("user_name")} ({user_info.get("user_type", "Unknown")})')
        header_cell.value = ' | '.join(header_parts)
        header_cell.font = Font(size=10, italic=True)
        header_cell.alignment = Alignment(horizontal='left')
        current_row += 2
    
    # Get best practices for this pilar
    # Normalize pilar name to handle both camelCase and snake_case
    pilar_normalized = pilar.replace('_', '').lower()
    best_practices = evaluation_data.get('bestPractices', [])
    filtered_bps = [
        bp for bp in best_practices 
        if bp.get('pilar', '').replace('_', '').lower() == pilar_normalized
    ]
    
    # Group BPs by question (e.g., SEC01, SEC02, etc.)
    grouped_bps = {}
    for bp in filtered_bps:
        bp_id = bp.get('bpId', '')
        # Extract question prefix (e.g., "SEC01" from "SEC01-BP01")
        question_id = bp_id.split('-')[0] if '-' in bp_id else bp_id[:5]
        if question_id not in grouped_bps:
            grouped_bps[question_id] = []
        grouped_bps[question_id].append(bp)
    
    # Sort question IDs
    sorted_questions = sorted(grouped_bps.keys())
    
    # Process each question group
    for question_id in sorted_questions:
        bps_in_question = grouped_bps[question_id]
        
        # Calculate question score using same formula as evaluator
        # COMPLIANT=100, PENDING_REVIEW=50, NON_COMPLIANT=0
        score_map = {'COMPLIANT': 100.0, 'NON_COMPLIANT': 0.0, 'PENDING_REVIEW': 50.0}
        total_score = sum(score_map.get(bp.get('status', ''), 0.0) for bp in bps_in_question)
        total_count = len(bps_in_question)
        question_score = round(total_score / total_count, 1) if total_count > 0 else 0
        
        # Question header with percentage
        ws.merge_cells(f'A{current_row}:H{current_row}')
        question_cell = ws.cell(row=current_row, column=1)
        
        # Get question title from first BP
        question_title = bps_in_question[0].get('questionTitle', question_id)
        question_cell.value = f'{question_id}: {question_title} - {question_score}%'
        question_cell.font = Font(size=12, bold=True, color='FFFFFF')
        question_cell.fill = PatternFill(start_color=GF_RED, end_color=GF_RED, fill_type='solid')
        question_cell.alignment = Alignment(horizontal='left', vertical='center')
        
        current_row += 1
        
        # Column headers for this group
        headers = ['BP ID', 'Descripción', 'Estado', 'Severidad', 'Hallazgo', 'Evidencia', 'Remediación', 'Riesgo']
        for col, header in enumerate(headers, start=1):
            cell = ws.cell(row=current_row, column=col, value=header)
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill(start_color=GF_ORANGE, end_color=GF_ORANGE, fill_type='solid')
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        
        current_row += 1
        
        # Long-text columns that should wrap and align to the top (Descripción B,
        # Hallazgo E, Evidencia F, Remediación G, Riesgo H).
        wrap_cols = {2, 5, 6, 7, 8}
        wrap_alignment = Alignment(wrap_text=True, vertical='top')
        
        # Add BPs for this question
        for bp in bps_in_question:
            ws.cell(row=current_row, column=1, value=bp.get('bpId', ''))
            ws.cell(row=current_row, column=2, value=bp.get('description', ''))
            
            # Status with color coding
            status_cell = ws.cell(row=current_row, column=3, value=bp.get('status', ''))
            apply_status_color(status_cell, bp.get('status', ''))
            
            ws.cell(row=current_row, column=4, value=bp.get('severity', ''))
            ws.cell(row=current_row, column=5, value=bp.get('finding', ''))
            ws.cell(row=current_row, column=6, value=bp.get('evidence', ''))
            ws.cell(row=current_row, column=7, value=bp.get('remediation', ''))
            ws.cell(row=current_row, column=8, value=bp.get('risk', ''))
            
            # Apply borders and wrap on long-text columns
            for col in range(1, 9):
                cell = ws.cell(row=current_row, column=col)
                if col in wrap_cols:
                    cell.alignment = wrap_alignment
                cell.border = Border(
                    left=Side(style='thin'),
                    right=Side(style='thin'),
                    top=Side(style='thin'),
                    bottom=Side(style='thin')
                )
            
            current_row += 1
        
        # Add spacing between question groups
        current_row += 1
    
    # Add summary section at the end
    current_row += 2
    add_pilar_summary(ws, filtered_bps, current_row)
    
    # Auto-adjust column widths
    ws.column_dimensions['A'].width = 15
    ws.column_dimensions['B'].width = 40
    ws.column_dimensions['C'].width = 15
    ws.column_dimensions['D'].width = 12
    ws.column_dimensions['E'].width = 50
    ws.column_dimensions['F'].width = 80
    ws.column_dimensions['G'].width = 50
    ws.column_dimensions['H'].width = 50


def add_pilar_summary(ws, bps, start_row):
    """Add summary statistics at the end of pilar sheet"""
    
    # Title
    ws.merge_cells(f'A{start_row}:H{start_row}')
    title_cell = ws.cell(row=start_row, column=1, value='RESUMEN DEL PILAR')
    title_cell.font = Font(size=14, bold=True, color='FFFFFF')
    title_cell.fill = PatternFill(start_color=GF_RED, end_color=GF_RED, fill_type='solid')
    title_cell.alignment = Alignment(horizontal='center', vertical='center')
    
    start_row += 2
    
    # Calculate statistics
    total_bps = len(bps)
    compliant = sum(1 for bp in bps if bp.get('status') == 'COMPLIANT')
    non_compliant = sum(1 for bp in bps if bp.get('status') == 'NON_COMPLIANT')
    pending = sum(1 for bp in bps if bp.get('status') == 'PENDING_REVIEW')
    
    # Calculate percentage using same formula as evaluator
    # COMPLIANT=100, PENDING_REVIEW=50, NON_COMPLIANT=0
    score_map = {'COMPLIANT': 100.0, 'NON_COMPLIANT': 0.0, 'PENDING_REVIEW': 50.0}
    total_score = sum(score_map.get(bp.get('status', ''), 0.0) for bp in bps)
    percentage = round(total_score / total_bps, 1) if total_bps > 0 else 0
    
    # Overall statistics
    ws.cell(row=start_row, column=1, value='Total BPs Evaluadas:')
    ws.cell(row=start_row, column=1).font = Font(bold=True)
    ws.cell(row=start_row, column=2, value=total_bps)
    
    ws.cell(row=start_row + 1, column=1, value='% Cumplimiento:')
    ws.cell(row=start_row + 1, column=1).font = Font(bold=True)
    ws.cell(row=start_row + 1, column=2, value=f'{percentage}%')
    
    ws.cell(row=start_row + 2, column=1, value='BPs CUMPLEN:')
    ws.cell(row=start_row + 2, column=1).font = Font(bold=True)
    ws.cell(row=start_row + 2, column=2, value=compliant)
    
    ws.cell(row=start_row + 3, column=1, value='BPs NO CUMPLEN:')
    ws.cell(row=start_row + 3, column=1).font = Font(bold=True)
    ws.cell(row=start_row + 3, column=2, value=non_compliant)
    
    ws.cell(row=start_row + 4, column=1, value='BPs PENDIENTES:')
    ws.cell(row=start_row + 4, column=1).font = Font(bold=True)
    ws.cell(row=start_row + 4, column=2, value=pending)
    
    # Severity breakdown by status
    start_row += 6
    ws.cell(row=start_row, column=1, value='DESGLOSE POR SEVERIDAD')
    ws.cell(row=start_row, column=1).font = Font(size=12, bold=True)
    start_row += 1
    
    # Headers
    ws.cell(row=start_row, column=1, value='Estado')
    ws.cell(row=start_row, column=2, value='Crítica')
    ws.cell(row=start_row, column=3, value='Alta')
    ws.cell(row=start_row, column=4, value='Media')
    ws.cell(row=start_row, column=5, value='Baja')
    for col in range(1, 6):
        ws.cell(row=start_row, column=col).font = Font(bold=True)
        ws.cell(row=start_row, column=col).fill = PatternFill(start_color=GF_ORANGE, end_color=GF_ORANGE, fill_type='solid')
        ws.cell(row=start_row, column=col).font = Font(bold=True, color='FFFFFF')
    
    start_row += 1
    
    # Calculate severity breakdown
    statuses = ['COMPLIANT', 'NON_COMPLIANT', 'PENDING_REVIEW']
    status_labels = {'COMPLIANT': 'Cumplen', 'NON_COMPLIANT': 'No Cumplen', 'PENDING_REVIEW': 'Pendientes'}
    severities = ['CRITICAL', 'HIGH', 'MEDIUM', 'LOW']
    
    for status in statuses:
        ws.cell(row=start_row, column=1, value=status_labels.get(status, status))
        ws.cell(row=start_row, column=1).font = Font(bold=True)
        
        for idx, severity in enumerate(severities, start=2):
            count = sum(1 for bp in bps if bp.get('status') == status and bp.get('severity') == severity)
            ws.cell(row=start_row, column=idx, value=count)
        
        start_row += 1


def apply_header_style(ws, row, num_cols):
    """Apply header styling to a row"""
    for col in range(1, num_cols + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill(start_color=GF_ORANGE, end_color=GF_ORANGE, fill_type='solid')
        cell.alignment = Alignment(horizontal='center', vertical='center')


def apply_status_color(cell, status):
    """Apply color coding based on status"""
    if status == 'COMPLIANT':
        cell.fill = PatternFill(start_color='C6EFCE', end_color='C6EFCE', fill_type='solid')
        cell.font = Font(color='006100')
    elif status == 'NON_COMPLIANT':
        cell.fill = PatternFill(start_color='FFC7CE', end_color='FFC7CE', fill_type='solid')
        cell.font = Font(color='9C0006')
    elif status == 'PENDING_REVIEW':
        cell.fill = PatternFill(start_color='FFEB9C', end_color='FFEB9C', fill_type='solid')
        cell.font = Font(color='9C6500')
