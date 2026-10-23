"""
Executive Report Generator Lambda
Generates business-oriented reports in DOCX format
"""

import json
import logging
import io
import os
from datetime import datetime
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from autowar_common.report_storage import upload_report, generate_signed_url

logger = logging.getLogger()
logger.setLevel(logging.INFO)

# GFourmis corporate palette
GF_ORANGE = RGBColor(0xF0, 0x60, 0x00)  # primary orange
GF_RED = RGBColor(0xC0, 0x10, 0x10)     # corporate red
GF_BLACK = RGBColor(0x1A, 0x1A, 0x1A)   # text black


def resolve_asset_path(filename):
    """Resolve a branding asset path with a Lambda-first fallback.

    Tries the Lambda layer mount (/opt/python/autowar_assets/<filename>) first,
    then falls back to the project-relative path computed from this handler's
    location so it also works in local tests. Returns the first existing path
    or None.
    """
    lambda_path = os.path.join('/opt/python', 'autowar_assets', filename)
    if os.path.exists(lambda_path):
        return lambda_path

    handler_dir = os.path.dirname(os.path.abspath(__file__))
    local_path = os.path.abspath(os.path.join(
        handler_dir, '..', 'layers', 'reports', 'python', 'autowar_assets', filename
    ))
    if os.path.exists(local_path):
        return local_path

    logger.info(f"Branding asset not found: {filename} (tried {lambda_path} and {local_path})")
    return None


# Ordered keyword rules mapping a best practice to a business risk scenario.
# Source of truth: .agents/tasks/riesgos_negocio.md (paraphrased, Spanish,
# managerial language). First rule with any keyword present in the haystack
# (bpId + description + finding, lowercased) wins.
RISK_RULES = [
    (
        ['mfa', 'multi-factor', 'sin mfa', 'root'],
        {
            'amenaza': 'Toma de control de la cuenta AWS.',
            'actor': 'Atacante con credenciales robadas (phishing o fuga de secretos).',
            'consecuencia': 'Control total de la nube del cliente: creación de recursos para '
                            'minería o fraude, borrado masivo de información y acceso a todos los '
                            'datos. Impacto crítico en continuidad, costos y confidencialidad.',
        },
    ),
    (
        ['s3', 'bucket', 'publico', 'público', 'public access', 'sin cifrado', 'encryption'],
        {
            'amenaza': 'Exfiltración (robo) de información sensible o exposición pública de datos.',
            'actor': 'Cualquiera en Internet (bots que escanean buckets abiertos) o un insider.',
            'consecuencia': 'Robo y publicación de datos de clientes y del negocio, con multas por '
                            'protección de datos (Ley 1581 de Habeas Data / SIC en Colombia), daño '
                            'reputacional y riesgo de extorsión con la información robada.',
        },
    ),
    (
        ['security group', 'puerto', '0.0.0.0', 'ssh', 'rdp', '22', '3389', 'ingress', 'abierto'],
        {
            'amenaza': 'Acceso no autorizado a servidores expuestos a Internet.',
            'actor': 'Atacante que escanea puertos abiertos y fuerza credenciales (fuerza bruta).',
            'consecuencia': 'Intrusión en los servidores, instalación de malware y pivote hacia la '
                            'red interna, con robo o cifrado de datos. Es el punto de entrada típico '
                            'de un ataque de ransomware (secuestro de información).',
        },
    ),
    (
        ['cloudtrail', 'log', 'logging', 'sin registro', 'auditoria', 'auditoría',
         'guardduty', 'deteccion', 'detección'],
        {
            'amenaza': 'Actividad maliciosa invisible: resulta imposible detectar o investigar un '
                       'incidente.',
            'actor': 'Atacante que opera sin dejar rastro o un insider que borra evidencias.',
            'consecuencia': 'Los ataques pasan desapercibidos durante meses, no hay forense tras una '
                            'brecha y se incumplen requisitos de auditoría y de respuesta a incidentes.',
        },
    ),
    (
        ['kms', 'cifrado', 'encryption', 'en reposo', 'en transito', 'en tránsito',
         'tls', 'certificado', 'acm'],
        {
            'amenaza': 'Interceptación o lectura de datos en tránsito o en reposo.',
            'actor': 'Atacante en la red (man-in-the-middle) o con acceso a discos y backups.',
            'consecuencia': 'Robo de datos confidenciales, credenciales y secretos; suplantación de '
                            'servicios e incumplimiento regulatorio por no cifrar datos personales '
                            '(Ley 1581 de Habeas Data / SIC).',
        },
    ),
    (
        ['waf', 'cloudfront', 'frontend', 'web', 'inyeccion', 'inyección', 'xss', 'sql'],
        {
            'amenaza': 'Suplantación del frontend, defacement, inyección de código y robo de sesiones.',
            'actor': 'Atacante web automatizado o dirigido.',
            'consecuencia': 'Suplantación de la aplicación para engañar a los usuarios (phishing con '
                            'su marca), robo de credenciales de clientes, manipulación de contenido '
                            'y caída del sitio.',
        },
    ),
    (
        ['iam', 'permiso', 'privilegio', 'politica', 'política', '*:*', 'admin', 'acceso'],
        {
            'amenaza': 'Abuso de permisos excesivos y escalada de privilegios.',
            'actor': 'Insider malicioso, cuenta comprometida o un error que un atacante aprovecha.',
            'consecuencia': 'Un único acceso comprometido se convierte en control total, con '
                            'movimiento lateral y acceso a datos y recursos más allá de lo necesario.',
        },
    ),
    (
        ['backup', 'snapshot', 'pitr', 'recuperacion', 'recuperación', 'retention', 'retención', 'dr'],
        {
            'amenaza': 'Pérdida irreversible de datos e incapacidad de recuperarse de un incidente.',
            'actor': 'Ransomware que cifra datos y backups, o un borrado accidental o malicioso.',
            'consecuencia': 'Pérdida permanente de información del negocio, parada prolongada de la '
                            'operación y pago de rescate sin garantía de recuperación.',
        },
    ),
    (
        ['multi-az', 'alta disponibilidad', 'single point', 'failover', 'autoscaling', 'auto scaling'],
        {
            'amenaza': 'Caída total del servicio ante el fallo de un componente o una zona.',
            'actor': 'Fallo de infraestructura, pico de demanda o un ataque de denegación (DoS).',
            'consecuencia': 'Indisponibilidad del servicio, pérdida de ventas y operación, '
                            'incumplimiento de SLA y daño reputacional.',
        },
    ),
    (
        ['vulnerab', 'parche', 'patch', 'cve', 'imagen', 'hardening', 'inspector'],
        {
            'amenaza': 'Explotación de vulnerabilidades conocidas sin parchar.',
            'actor': 'Atacante que usa exploits públicos contra software desactualizado.',
            'consecuencia': 'Compromiso de servidores, ransomware y robo de datos; es el vector más '
                            'común de brechas masivas.',
        },
    ),
]

# Fallback scenarios by pilar (normalized key) when no keyword rule matches.
RISK_FALLBACK_BY_PILAR = {
    'security': {
        'amenaza': 'Exposición de datos y sistemas a accesos no autorizados.',
        'actor': 'Actor malicioso externo o interno.',
        'consecuencia': 'Riesgo de robo o secuestro de información, suplantación y fraude.',
    },
    'reliability': {
        'amenaza': 'Interrupciones del servicio y pérdida de datos ante fallos.',
        'actor': 'Fallo de infraestructura o de componentes de la carga de trabajo.',
        'consecuencia': 'Impacto en la continuidad del negocio y en la experiencia del cliente.',
    },
    'performance': {
        'amenaza': 'Degradación de la experiencia por recursos mal dimensionados.',
        'actor': 'Picos de demanda o configuración ineficiente.',
        'consecuencia': 'Mala experiencia de usuario y sobrecostos operativos.',
    },
    'cost': {
        'amenaza': 'Fuga de dinero por recursos ociosos o sobredimensionados.',
        'actor': 'Configuración ineficiente y falta de gobierno de costos.',
        'consecuencia': 'Gasto innecesario que erosiona el margen del negocio.',
    },
    'operationalExcellence': {
        'amenaza': 'Incidentes mal gestionados y cambios riesgosos.',
        'actor': 'Procesos manuales y falta de visibilidad operativa.',
        'consecuencia': 'Mayor tiempo de resolución, errores recurrentes y baja previsibilidad.',
    },
    'sustainability': {
        'amenaza': 'Uso ineficiente de recursos en la nube.',
        'actor': 'Dimensionamiento y patrones de consumo subóptimos.',
        'consecuencia': 'Impacto ambiental y de costo evitable.',
    },
}


def map_risk_scenarios(bp):
    """Map a best practice to a business risk scenario.

    Returns a dict {'amenaza', 'actor', 'consecuencia'} based on keyword rules
    (searched across bpId + description + finding, case-insensitive). If no rule
    matches, falls back to the pilar-level generic scenario (defaulting to
    'security' as a last resort for unknown pilares).
    """
    haystack = f"{bp.get('bpId', '')} {bp.get('description', '')} {bp.get('finding', '')}".lower()

    for keywords, scenario in RISK_RULES:
        if any(kw in haystack for kw in keywords):
            return scenario

    # Fallback by pilar
    pilar_normalized = bp.get('pilar', '').replace('_', '').lower()
    pilar_key_map = {
        'security': 'security',
        'reliability': 'reliability',
        'performance': 'performance',
        'cost': 'cost',
        'operationalexcellence': 'operationalExcellence',
        'sustainability': 'sustainability',
    }
    key = pilar_key_map.get(pilar_normalized, 'security')
    return RISK_FALLBACK_BY_PILAR[key]


def lambda_handler(event, context):
    """
    Generate executive report in DOCX format
    
    Input:
    {
        "accountId": "...",
        "evaluationTimestamp": "...",
        "evaluationData": {...},
        "selectedPilares": ["security", "reliability", ...]
    }
    
    Output:
    {
        "s3Key": "...",
        "signedUrl": "...",
        "expiresIn": 3600
    }
    """
    try:
        # Parse input
        if isinstance(event.get('body'), str):
            body = json.loads(event['body'])
        else:
            body = event
        
        account_id = body.get('accountId')
        evaluation_timestamp = body.get('evaluationTimestamp')
        evaluation_data = body.get('evaluationData')
        selected_pilares = body.get('selectedPilares', [])
        user_info = body.get('userInfo', {})
        
        logger.info(f"Generating executive report for account {account_id}")
        
        # Generate DOCX report
        docx_content = generate_executive_report(
            evaluation_data,
            selected_pilares,
            account_id,
            evaluation_timestamp,
            user_info
        )
        
        # Upload to S3
        s3_key = upload_report(
            account_id,
            evaluation_timestamp,
            'executive',
            docx_content,
            'docx'
        )
        
        # Generate signed URL
        signed_url = generate_signed_url(s3_key, expiration_seconds=3600)
        
        logger.info(f"Executive report generated: {s3_key}")
        
        return {
            'statusCode': 200,
            'headers': {
                'Content-Type': 'application/json',
                'Access-Control-Allow-Origin': '*'
            },
            'body': json.dumps({
                's3Key': s3_key,
                'signedUrl': signed_url,
                'expiresIn': 3600
            })
        }
        
    except Exception as e:
        logger.error(f"Error generating executive report: {str(e)}", exc_info=True)
        return {
            'statusCode': 500,
            'headers': {
                'Content-Type': 'application/json',
                'Access-Control-Allow-Origin': '*'
            },
            'body': json.dumps({
                'error': 'Failed to generate executive report',
                'message': str(e)
            })
        }


def generate_executive_report(evaluation_data, selected_pilares, account_id, evaluation_timestamp, user_info=None):
    """Generate executive report DOCX content"""
    
    doc = Document()
    
    # Title page
    add_title_page(doc, account_id, evaluation_timestamp, user_info)
    
    # Executive summary
    add_executive_summary(doc, evaluation_data, selected_pilares)
    
    # Pillar analysis sections
    for pilar in selected_pilares:
        add_pilar_section(doc, evaluation_data, pilar)
    
    # Conclusions and recommendations
    add_conclusions(doc, evaluation_data, selected_pilares)
    
    # Save to bytes
    output = io.BytesIO()
    doc.save(output)
    output.seek(0)
    
    return output.read()


def _color_heading(heading, color):
    """Apply a font color to all runs of a heading (defensive if it has none)."""
    for run in heading.runs:
        run.font.color.rgb = color


def add_title_page(doc, account_id, evaluation_timestamp, user_info=None):
    """Add title page"""
    
    # GFourmis logo centered at the top
    logo_path = resolve_asset_path('gfourmis-logo.png')
    if logo_path:
        try:
            doc.add_picture(logo_path, width=Inches(2.5))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        except Exception as e:
            logger.info(f"Could not embed logo in title page: {e}")
    
    # Title
    title = doc.add_heading('Reporte Ejecutivo', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _color_heading(title, GF_ORANGE)
    
    subtitle = doc.add_heading('AWS Well-Architected Review', level=2)
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Brand footer line
    brand = doc.add_paragraph()
    brand.alignment = WD_ALIGN_PARAGRAPH.CENTER
    brand_run = brand.add_run('GFourmis — Strategic & digital solutions')
    brand_run.font.size = Pt(11)
    brand_run.font.color.rgb = GF_RED
    brand_run.italic = True
    
    # Metadata
    doc.add_paragraph()
    doc.add_paragraph()
    
    metadata = doc.add_paragraph()
    metadata.add_run(f'Cuenta AWS: ').bold = True
    metadata.add_run(account_id)
    
    metadata = doc.add_paragraph()
    metadata.add_run(f'Fecha de Evaluación: ').bold = True
    metadata.add_run(evaluation_timestamp)
    
    metadata = doc.add_paragraph()
    metadata.add_run(f'Fecha de Generación: ').bold = True
    metadata.add_run(datetime.utcnow().strftime('%Y-%m-%d %H:%M:%S UTC'))
    
    # User information
    if user_info:
        metadata = doc.add_paragraph()
        metadata.add_run(f'Usuario: ').bold = True
        metadata.add_run(f"{user_info.get('user_name', 'Unknown')} ({user_info.get('user_type', 'Unknown')})")
        
        metadata = doc.add_paragraph()
        metadata.add_run(f'User ID: ').bold = True
        metadata.add_run(user_info.get('user_id', 'Unknown'))
    
    # Page break
    doc.add_page_break()


def add_executive_summary(doc, evaluation_data, selected_pilares):
    """Add executive summary section"""
    
    _color_heading(doc.add_heading('Resumen Ejecutivo', level=1), GF_ORANGE)
    
    pillar_scores = evaluation_data.get('pillarScores', {})
    best_practices = evaluation_data.get('bestPractices', [])
    
    # Calculate overall metrics
    # Normalize pilar names for comparison
    selected_pilares_normalized = [p.replace('_', '').lower() for p in selected_pilares]
    filtered_bps = [
        bp for bp in best_practices 
        if bp.get('pilar', '').replace('_', '').lower() in selected_pilares_normalized
    ]
    total_bps = len(filtered_bps)
    non_compliant = sum(1 for bp in filtered_bps if bp.get('status') == 'NON_COMPLIANT')
    compliant = sum(1 for bp in filtered_bps if bp.get('status') == 'COMPLIANT')
    pending = sum(1 for bp in filtered_bps if bp.get('status') == 'PENDING_REVIEW')
    
    # Calculate overall score using same formula as evaluator
    score_map = {'COMPLIANT': 100.0, 'NON_COMPLIANT': 0.0, 'PENDING_REVIEW': 50.0}
    overall_score = round(sum(score_map.get(bp.get('status', ''), 0.0) for bp in filtered_bps) / total_bps, 1) if total_bps > 0 else 0
    
    # Overall assessment
    doc.add_paragraph(
        f'Este reporte presenta un análisis estratégico de la evaluación Well-Architected '
        f'realizada para la cuenta AWS {evaluation_data.get("metadata", {}).get("accountId", "N/A")}. '
        f'Se evaluaron {len(selected_pilares)} pilares con un total de {total_bps} mejores prácticas.'
    )
    
    doc.add_paragraph()
    
    # Key findings
    doc.add_heading('Hallazgos Clave', level=2)
    
    findings_para = doc.add_paragraph()
    findings_para.add_run(f'• Score Global: ').bold = True
    findings_para.add_run(f'{overall_score}%')
    
    findings_para = doc.add_paragraph()
    findings_para.add_run(f'• Prácticas Conformes: ').bold = True
    findings_para.add_run(f'{compliant} de {total_bps} ({round(compliant/total_bps*100, 1) if total_bps > 0 else 0}%)')
    
    findings_para = doc.add_paragraph()
    findings_para.add_run(f'• Prácticas No Conformes: ').bold = True
    findings_para.add_run(f'{non_compliant} de {total_bps} ({round(non_compliant/total_bps*100, 1) if total_bps > 0 else 0}%)')
    
    findings_para = doc.add_paragraph()
    findings_para.add_run(f'• Prácticas Pendientes de Revisión: ').bold = True
    findings_para.add_run(f'{pending} de {total_bps} ({round(pending/total_bps*100, 1) if total_bps > 0 else 0}%)')
    
    doc.add_paragraph()
    
    # Pillar scores
    doc.add_heading('Scores por Pilar', level=2)
    
    pillar_names = {
        'security': 'Seguridad',
        'reliability': 'Fiabilidad',
        'performance': 'Rendimiento',
        'cost': 'Optimización de Costos',
        'operationalExcellence': 'Excelencia Operacional',
        'sustainability': 'Sostenibilidad'
    }
    
    for pilar in selected_pilares:
        score = pillar_scores.get(pilar, 0)
        score_para = doc.add_paragraph()
        score_para.add_run(f'• {pillar_names.get(pilar, pilar)}: ').bold = True
        score_para.add_run(f'{score}/100')
        
        # Add assessment
        if score >= 80:
            score_para.add_run(' - Excelente')
        elif score >= 60:
            score_para.add_run(' - Bueno, con oportunidades de mejora')
        elif score >= 40:
            score_para.add_run(' - Requiere atención')
        else:
            score_para.add_run(' - Requiere acción inmediata')
    
    doc.add_page_break()


def add_pilar_section(doc, evaluation_data, pilar):
    """Add analysis section for a specific pilar"""
    
    pillar_names = {
        'security': 'Seguridad',
        'reliability': 'Fiabilidad',
        'performance': 'Rendimiento',
        'cost': 'Optimización de Costos',
        'operationalExcellence': 'Excelencia Operacional',
        'sustainability': 'Sostenibilidad'
    }
    
    pillar_descriptions = {
        'security': 'La seguridad es fundamental para proteger la información, sistemas y activos de su organización.',
        'reliability': 'La Fiabilidad asegura que su carga de trabajo funcione correctamente y de manera consistente.',
        'performance': 'El rendimiento eficiente permite utilizar los recursos de manera óptima.',
        'cost': 'La optimización de costos permite ejecutar sistemas para entregar valor de negocio al menor precio.',
        'operationalExcellence': 'La excelencia operacional permite ejecutar y monitorear sistemas para entregar valor de negocio.',
        'sustainability': 'La sostenibilidad minimiza los impactos ambientales de ejecutar cargas de trabajo en la nube.'
    }
    
    _color_heading(doc.add_heading(f'Análisis: {pillar_names.get(pilar, pilar)}', level=1), GF_ORANGE)
    
    # Description
    doc.add_paragraph(pillar_descriptions.get(pilar, ''))
    doc.add_paragraph()
    
    # Score
    pillar_scores = evaluation_data.get('pillarScores', {})
    score = pillar_scores.get(pilar, 0)
    
    score_para = doc.add_paragraph()
    score_para.add_run('Score del Pilar: ').bold = True
    score_para.add_run(f'{score}/100')
    
    doc.add_paragraph()
    
    # Key findings for this pilar
    doc.add_heading('Hallazgos Principales', level=2)
    
    best_practices = evaluation_data.get('bestPractices', [])
    # Normalize pilar name for comparison
    pilar_normalized = pilar.replace('_', '').lower()
    pilar_bps = [
        bp for bp in best_practices 
        if bp.get('pilar', '').replace('_', '').lower() == pilar_normalized
    ]
    
    # Count by status
    non_compliant_bps = [bp for bp in pilar_bps if bp.get('status') == 'NON_COMPLIANT']
    
    if non_compliant_bps:
        doc.add_paragraph(f'Se identificaron {len(non_compliant_bps)} áreas que requieren atención:')
        
        # Show top 5 non-compliant items
        for bp in non_compliant_bps[:5]:
            bp_para = doc.add_paragraph(style='List Bullet')
            bp_para.add_run(f'{bp.get("bpId", "")}: ').bold = True
            bp_para.add_run(bp.get('finding', ''))
    else:
        doc.add_paragraph('Este pilar cumple con todas las mejores prácticas evaluadas.')
    
    doc.add_paragraph()
    
    # Risk scenarios and consequences
    add_risk_scenarios_subsection(doc, non_compliant_bps)
    
    # Business impact
    doc.add_heading('Impacto de Negocio', level=2)
    
    if score < 60:
        doc.add_paragraph(
            f'El score actual de {score}/100 indica riesgos significativos que pueden afectar '
            f'la continuidad del negocio, la seguridad de los datos, o los costos operacionales. '
            f'Se recomienda priorizar las remediaciones identificadas.'
        )
    elif score < 80:
        doc.add_paragraph(
            f'El score de {score}/100 es aceptable, pero existen oportunidades de mejora que '
            f'pueden optimizar la operación y reducir riesgos potenciales.'
        )
    else:
        doc.add_paragraph(
            f'El score de {score}/100 indica un excelente nivel de madurez en este pilar. '
            f'Se recomienda mantener las prácticas actuales y revisar periódicamente.'
        )
    
    doc.add_page_break()


def add_risk_scenarios_subsection(doc, non_compliant_bps):
    """Add the 'Escenarios de Riesgo y Consecuencias' subsection for a pilar.

    For each NON_COMPLIANT best practice with CRITICAL or HIGH severity (criticals
    first, limited to the first ~8), writes a readable block describing the
    threat, actor and business consequence in managerial Spanish.
    """
    doc.add_heading('Escenarios de Riesgo y Consecuencias', level=2)

    severity_order = {'CRITICAL': 0, 'HIGH': 1}
    high_risk_bps = [
        bp for bp in non_compliant_bps
        if bp.get('severity') in severity_order
    ]
    high_risk_bps.sort(key=lambda bp: severity_order.get(bp.get('severity'), 99))
    high_risk_bps = high_risk_bps[:8]

    if not high_risk_bps:
        doc.add_paragraph(
            'No se identificaron escenarios de riesgo alto (hallazgos críticos o altos no '
            'conformes) en este pilar.'
        )
        return

    for bp in high_risk_bps:
        scenario = map_risk_scenarios(bp)

        threat_para = doc.add_paragraph()
        threat_run = threat_para.add_run(f'{bp.get("bpId", "")} — {scenario["amenaza"]}')
        threat_run.bold = True

        actor_para = doc.add_paragraph()
        actor_para.add_run('Actor: ').bold = True
        actor_para.add_run(scenario['actor'])

        consequence_para = doc.add_paragraph()
        consequence_para.add_run('Consecuencia: ').bold = True
        consequence_para.add_run(scenario['consecuencia'])

        doc.add_paragraph()

    note_para = doc.add_paragraph()
    note_run = note_para.add_run(
        'Nota: estos escenarios describen el riesgo potencial si el hallazgo no se remedia; '
        'no implican que el ataque haya ocurrido.'
    )
    note_run.italic = True

    doc.add_paragraph()


def add_conclusions(doc, evaluation_data, selected_pilares):
    """Add conclusions and recommendations section"""
    
    _color_heading(doc.add_heading('Conclusiones y Recomendaciones', level=1), GF_ORANGE)
    
    pillar_scores = evaluation_data.get('pillarScores', {})
    
    # Calculate average score
    scores = [pillar_scores.get(p, 0) for p in selected_pilares]
    avg_score = sum(scores) / len(scores) if scores else 0
    
    doc.add_paragraph(
        f'La evaluación Well-Architected ha identificado el estado actual de la arquitectura '
        f'con un score promedio de {round(avg_score, 1)}/100 en los pilares evaluados.'
    )
    
    doc.add_paragraph()
    
    # Priority recommendations
    doc.add_heading('Recomendaciones Prioritarias', level=2)
    
    # Find pillar with lowest score
    if scores:
        min_score = min(scores)
        min_pilar = selected_pilares[scores.index(min_score)]
        
        pillar_names = {
            'security': 'Seguridad',
            'reliability': 'Fiabilidad',
            'performance': 'Rendimiento',
            'cost': 'Optimización de Costos',
            'operationalExcellence': 'Excelencia Operacional',
            'sustainability': 'Sostenibilidad'
        }
        
        doc.add_paragraph(
            f'1. Priorizar remediaciones en el pilar de {pillar_names.get(min_pilar, min_pilar)} '
            f'(score: {min_score}/100), ya que presenta el mayor número de oportunidades de mejora.'
        )
    
    doc.add_paragraph(
        '2. Establecer un plan de remediación con plazos definidos para cada hallazgo identificado.'
    )
    
    doc.add_paragraph(
        '3. Implementar revisiones periódicas (trimestrales) para monitorear el progreso y '
        'mantener la arquitectura alineada con las mejores prácticas.'
    )
    
    doc.add_paragraph(
        '4. Considerar la implementación de automatización para mantener el cumplimiento continuo '
        'de las mejores prácticas.'
    )
    
    doc.add_paragraph()
    
    # Next steps
    doc.add_heading('Próximos Pasos', level=2)
    
    doc.add_paragraph(
        '• Revisar el reporte técnico detallado para obtener información específica sobre cada hallazgo'
    )
    
    doc.add_paragraph(
        '• Asignar responsables y plazos para cada remediación prioritaria'
    )
    
    doc.add_paragraph(
        '• Establecer métricas de seguimiento para medir el progreso'
    )
    
    doc.add_paragraph(
        '• Programar la próxima evaluación Well-Architected'
    )
